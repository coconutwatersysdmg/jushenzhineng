# -*- coding: utf-8 -*-
"""单表：现场联调所需全部字段"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(r"E:\Homework\AA定位\jushenzhineng_v3\docs\现场所需字段表.docx")


def font(run, name="微软雅黑", size=10, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def cell(c, text, bold=False, size=9):
    c.text = ""
    run = c.paragraphs[0].add_run(str(text))
    font(run, "微软雅黑", size, bold)


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(1.8)
        s.bottom_margin = Cm(1.8)
        s.left_margin = Cm(1.8)
        s.right_margin = Cm(1.8)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("具身智能装载 — 现场所需字段表")
    font(r, "微软雅黑", 14, True)

    headers = ["分类", "字段", "类型", "说明", "默认/示例", "配置位置"]
    rows = [
        # 装载计划
        ["装载计划", "cargo_code", "string", "货物编码", "CARGO-001", "data\\loading_plan.json → items[]"],
        ["装载计划", "cargo_name", "string", "货物名称", "示例货物", "同上"],
        ["装载计划", "quantity", "int", "件数（按件展开队列）", "3", "同上"],
        ["装载计划", "length_mm", "number", "货长 mm", "1200", "同上"],
        ["装载计划", "width_mm", "number", "货宽 mm", "1000", "同上"],
        ["装载计划", "height_mm", "number", "货高 mm", "900", "同上"],
        ["装载计划", "pallet_reference_width_mm", "number", "托盘参考宽度 mm", "1200", "同上"],
        # PLC
        ["PLC", "ip", "string", "PLC IP 地址", "192.168.6.6", "config\\external_devices_config.py → PLC"],
        ["PLC", "port", "int", "Modbus TCP 端口", "502", "同上"],
        ["PLC", "ack_timeout_ms", "int", "ACK 超时 ms", "5000", "同上"],
        # Livox
        ["Livox雷达", "host_ip", "string", "连雷达网卡 IPv4（须=ipconfig）", "192.168.1.50", "LIVOX + mid360s_config.json"],
        ["Livox雷达", "capture_ms", "int", "采点时长 ms", "3000", "LIVOX"],
        ["Livox雷达", "save_dir", "string", "点云保存目录", "data\\lidar", "LIVOX"],
        ["Livox雷达", "max_points", "int", "最大点数", "300000", "LIVOX"],
        ["Livox雷达", "timeout_sec", "int", "采集超时 s", "20", "LIVOX"],
        ["Livox雷达", "lidar_ports.cmd_data_port", "int", "雷达 cmd 端口", "56100", "LIVOX / mid360s_config.json"],
        ["Livox雷达", "lidar_ports.push_msg_port", "int", "雷达 push 端口", "56200", "同上"],
        ["Livox雷达", "lidar_ports.point_data_port", "int", "雷达点云端口", "56300", "同上"],
        ["Livox雷达", "lidar_ports.imu_data_port", "int", "雷达 IMU 端口", "56400", "同上"],
        ["Livox雷达", "lidar_ports.log_data_port", "int", "雷达 log 端口", "56500", "同上"],
        ["Livox雷达", "host_ports.cmd_data_port", "int", "主机 cmd 端口", "56101", "同上"],
        ["Livox雷达", "host_ports.push_msg_port", "int", "主机 push 端口", "56201", "同上"],
        ["Livox雷达", "host_ports.point_data_port", "int", "主机点云端口", "56301", "同上"],
        ["Livox雷达", "host_ports.imu_data_port", "int", "主机 IMU 端口", "56401", "同上"],
        ["Livox雷达", "host_ports.log_data_port", "int", "主机 log 端口", "56501", "同上"],
        ["Livox雷达", "mid360_json_path", "string", "雷达 JSON 配置路径", "third_party\\livox_runtime\\mid360s_config.json", "LIVOX"],
        ["Livox雷达", "exe_path", "string", "采集程序路径", "third_party\\livox_runtime\\livox_realtime_select_and_move.exe", "LIVOX"],
        # Camera
        ["D435i相机", "backend", "string", "相机后端", "realsense", "CAMERA"],
        ["D435i相机", "serial", "string", "序列号；空=自动第一台", "（空）", "CAMERA"],
        ["D435i相机", "color_width", "int", "彩色宽", "1280", "CAMERA"],
        ["D435i相机", "color_height", "int", "彩色高", "720", "CAMERA"],
        ["D435i相机", "depth_width", "int", "深度宽", "1280", "CAMERA"],
        ["D435i相机", "depth_height", "int", "深度高", "720", "CAMERA"],
        ["D435i相机", "fps", "int", "帧率", "30", "CAMERA"],
        ["D435i相机", "depth_unit_mm", "number", "深度单位 mm", "1.0", "CAMERA"],
        ["D435i相机", "extrinsic_file", "string", "相机外参文件", "config\\camera_extrinsic.json", "CAMERA"],
        ["D435i相机", "intrinsic_file", "string", "相机内参文件", "config\\sensor_coordinate_config\\camera_intrinsic.json", "CAMERA"],
        ["D435i相机", "plc_reference_r_deg", "number", "标定参考 R 角 °", "-80.0", "CAMERA / camera_extrinsic.json"],
        ["D435i相机", "capture_dir", "string", "抓图目录", "workdir\\camera_captures", "CAMERA"],
        ["D435i相机", "物理连接", "—", "USB3 + RealSense 系统驱动", "—", "硬件/系统"],
    ]

    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, ht in enumerate(headers):
        cell(t.rows[0].cells[i], ht, bold=True, size=9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell(t.rows[ri + 1].cells[ci], val, size=8.5)

    # prefer primary path; if locked try nothing else needed
    try:
        doc.save(str(OUT))
        print("OK", OUT, OUT.stat().st_size)
    except PermissionError:
        alt = OUT.with_name("现场所需字段表_新.docx")
        doc.save(str(alt))
        print("OK", alt, alt.stat().st_size)


if __name__ == "__main__":
    main()
