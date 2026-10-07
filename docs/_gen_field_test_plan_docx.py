# -*- coding: utf-8 -*-
"""内部备忘：三模块 — 输入字段 / 三外设 / 总体流程"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


PATHS = [
    Path(r"E:\Homework\AA定位\jushenzhineng_v3\docs\现场联调备忘.docx"),
    Path(r"E:\Homework\AA定位\jushenzhineng_v3\docs\现场软件测试计划.docx"),
]


def font(run, name="微软雅黑", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        font(r, "微软雅黑", 14 if level == 1 else 12, True)


def p(doc, text, size=11, bold=False):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.line_spacing = 1.25
    run = para.add_run(text)
    font(run, "微软雅黑", size, bold)


def cell(c, text, bold=False, size=9.5):
    c.text = ""
    run = c.paragraphs[0].add_run(text)
    font(run, "微软雅黑", size, bold)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, ht in enumerate(headers):
        cell(t.rows[0].cells[i], ht, bold=True)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell(t.rows[ri + 1].cells[ci], str(val))
    doc.add_paragraph()


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("具身智能装载 — 现场联调备忘")
    font(r, "微软雅黑", 16, True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("jushenzhineng_v3 · field 模式 · 内部")
    font(r, "微软雅黑", 10)

    # ---------- 1 ----------
    h(doc, "1. 现场输入数据说明", 1)
    p(doc, "不管立库用 HTTP、文件，还是别的方式下发，最终都要落到下面字段。少一列现场就对不齐。")
    table(
        doc,
        ["字段", "类型", "说明", "示例"],
        [
            ["cargo_code", "string", "货物编码", "CARGO-001"],
            ["cargo_name", "string", "货物名称", "示例货物"],
            ["quantity", "int", "件数（按件展开成装载队列）", "3"],
            ["length_mm", "number", "货长，单位 mm", "1200"],
            ["width_mm", "number", "货宽，单位 mm", "1000"],
            ["height_mm", "number", "货高，单位 mm", "900"],
            ["pallet_reference_width_mm", "number", "托盘参考宽度，单位 mm", "1200"],
        ],
    )
    p(doc, "当前软件读法：本地 JSON 文件 data\\loading_plan.json，结构为 { \"items\": [ {上述字段…}, … ] }。")
    p(doc, "若改成立库 HTTP：地址/鉴权按现场协议配；响应里把各字段映射到上表即可。协议细节对接时再填，字段集合不变。")

    # ---------- 2 ----------
    h(doc, "2. 三个外接设备连接所需信息", 1)
    p(doc, "配置入口：config\\external_devices_config.py。雷达 JSON 与 host_ip 保持一致。")

    h(doc, "2.1 PLC（Modbus TCP）", 2)
    table(
        doc,
        ["需要的信息", "默认 / 现状", "备注"],
        [
            ["IP", "192.168.6.6", "改成现场 PLC 地址"],
            ["端口", "502", "一般不变"],
            ["ack 超时(ms)", "5000", "PLC[\"ack_timeout_ms\"]"],
            ["工控机↔PLC 网通", "—", "网线/网段/防火墙放行 502"],
        ],
    )

    h(doc, "2.2 Livox Mid-360（雷达）", 2)
    table(
        doc,
        ["需要的信息", "默认 / 现状", "备注"],
        [
            ["电脑 host_ip", "192.168.1.50", "必须等于连雷达网卡的 ipconfig"],
            ["采点时长(ms)", "3000", "LIVOX[\"capture_ms\"]"],
            ["点云保存目录", "data\\lidar", "LIVOX[\"save_dir\"]"],
            ["雷达端口", "56100/56200/56300/56400/56500", "cmd/push/point/imu/log"],
            ["主机端口", "56101/56201/56301/56401/56501", "与上面对应"],
            ["配置文件", "third_party\\livox_runtime\\mid360s_config.json", "host_ip 一并改"],
            ["供电+网线", "—", "独立网段，通了才能采点"],
        ],
    )

    h(doc, "2.3 RealSense D435i（相机）", 2)
    table(
        doc,
        ["需要的信息", "默认 / 现状", "备注"],
        [
            ["物理连接", "USB3", "无 IP；勿用 USB2"],
            ["系统驱动", "Intel RealSense 驱动", "官方 Viewer 能出图再开本软件"],
            ["分辨率/帧率", "1280×720 @30", "彩色与深度一致"],
            ["序列号 serial", "空=自动第一台", "多机时填写"],
            ["外参文件", "config\\camera_extrinsic.json", "D435i→臂端"],
            ["内参文件", "config\\sensor_coordinate_config\\camera_intrinsic.json", "—"],
            ["参考 R 角", "-80°", "plc_reference_r_deg"],
            ["抓图目录", "workdir\\camera_captures", "—"],
        ],
    )

    # ---------- 3 ----------
    h(doc, "3. 总体流程（简述）", 1)
    p(doc, "WORLD 坐标，单位 mm：X 右、Y 车前进、Z 上。界面选「完全真实 / field」。")
    p(doc, "首轮：", bold=True)
    p(doc, "读计划 → 设备检查（PLC/雷达/相机）→ 货托偏移 → "
           "并行（插孔识别插取 ∥ 雷达找车粗点）→ 规划拍照位 → 车尾/车头拍角点 → "
           "角点识别（可人工审）→ 转到 WORLD → 判板型并划装载区 → "
           "邻托补偿 → 锁定放置目标 → 放前监测 → 放货 → 放后检测与反馈 → 回左侧初始位。")
    p(doc, "第 2 轮起：", bold=True)
    p(doc, "货托偏移 → 插孔插取 → 邻托补偿 → 锁目标 → 放前监测 → 放货 → 放后检测与反馈 → 返回。"
           "不再做雷达找车和角点建图，复用首轮板型与空间。")
    p(doc, "联调习惯：软件出坐标，对方确认后再动臂；本步失败优先重试。", size=10.5)

    saved = []
    for path in PATHS:
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(path))
            saved.append(f"{path.name} ({path.stat().st_size})")
        except PermissionError:
            saved.append(f"{path.name} (占用未写)")
    print("saved:", "; ".join(saved))


if __name__ == "__main__":
    main()
